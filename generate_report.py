import csv
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─── Load CSV data ────────────────────────────────────────────────────────────
rows = []
with open("benchmark_results.csv") as f:
    reader = csv.DictReader(f)
    for r in reader:
        if r["GridSize"].strip():
            rows.append({
                "size"    : int(r["GridSize"]),
                "cells"   : int(r["TotalCells"]),
                "cpu"     : float(r["CPU_Time_ms"]),
                "glob"    : float(r["CUDA_Global_ms"]),
                "shrd"    : float(r["CUDA_Shared_ms"]),
                "sp_glob" : float(r["Speedup_Global"]),
                "sp_shrd" : float(r["Speedup_Shared"]),
            })

# ─── Helpers ──────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        bdr = OxmlElement(f"w:{side}")
        bdr.set(qn("w:val"),   "single")
        bdr.set(qn("w:sz"),    "4")
        bdr.set(qn("w:space"), "0")
        bdr.set(qn("w:color"), "BFBFBF")
        tcBorders.append(bdr)
    tcPr.append(tcBorders)

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.runs[0]
    run.font.color.rgb = RGBColor(0x1A, 0x53, 0x76) if level == 1 else RGBColor(0x1F, 0x6E, 0x5E)
    run.font.bold = True
    return p

def body(doc, text, bold_parts=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    return p

# ─── Create Document ──────────────────────────────────────────────────────────
doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.8)
    section.right_margin  = Cm(2.8)

# ─── TITLE ────────────────────────────────────────────────────────────────────
title = doc.add_heading("Conway's Game of Life", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(0x0D, 0x3B, 0x6E)
title.runs[0].font.size = Pt(22)

sub = doc.add_paragraph("High-Performance Computing Assignment Report")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(13)
sub.runs[0].font.bold = True
sub.runs[0].font.color.rgb = RGBColor(0x2C, 0x6E, 0x90)

info = doc.add_paragraph("CPU (2D Array) vs. CUDA GPU Parallel Implementation")
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.runs[0].font.size = Pt(11)
info.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

# ─── 1. INTRODUCTION ──────────────────────────────────────────────────────────
heading(doc, "1. Introduction")
body(doc, (
    "Conway's Game of Life is a well-known cellular automaton invented by mathematician John Conway in 1970. "
    "It simulates how cells live, die, and reproduce on a 2D grid based on four simple rules applied to every "
    "cell simultaneously. Despite its simple rules, the simulation produces highly complex and interesting patterns."
))
body(doc, (
    "In this assignment, we implement Conway's Game of Life using two methods:"
))
bullet(doc, "Method 1 — CPU: A standard sequential implementation using 2D arrays in C++.")
bullet(doc, "Method 2 — GPU: A parallel implementation using NVIDIA CUDA C++, running on the NVIDIA GeForce MX350 GPU.")
body(doc, (
    "We measure the execution time for 100 generations across six different grid sizes and compare the "
    "performance of both methods."
))

# ─── 2. GAME RULES ────────────────────────────────────────────────────────────
heading(doc, "2. Game of Life Rules")
body(doc, "Every cell is either alive (1) or dead (0). At each generation, every cell is updated using these rules:")
bullet(doc, "A live cell with fewer than 2 live neighbors dies (underpopulation).")
bullet(doc, "A live cell with 2 or 3 live neighbors stays alive.")
bullet(doc, "A live cell with more than 3 live neighbors dies (overpopulation).")
bullet(doc, "A dead cell with exactly 3 live neighbors becomes alive (reproduction).")
body(doc, (
    "The grid uses toroidal (wrap-around) boundary conditions, meaning the left edge connects to the right edge "
    "and the top edge connects to the bottom edge."
))

# ─── 3. HARDWARE ──────────────────────────────────────────────────────────────
heading(doc, "3. Hardware & Software Configuration")

hw_data = [
    ("Component",           "Details"),
    ("CPU",                 "Intel Core i5 (Laptop)"),
    ("GPU",                 "NVIDIA GeForce MX350"),
    ("GPU Memory",          "2 GB GDDR5 VRAM"),
    ("GPU Compute Capability", "6.1"),
    ("GPU Multiprocessors", "5 Streaming Multiprocessors (SMs)"),
    ("Max Threads / Block", "1024"),
    ("CUDA Version",        "12.4 (nvcc compiler)"),
    ("OS",                  "Linux (Ubuntu)"),
    ("Compiler Flags",      "-O3 -std=c++17 (maximum optimization)"),
    ("Iterations",          "100 generations per benchmark run"),
]

tbl = doc.add_table(rows=len(hw_data), cols=2)
tbl.style = "Table Grid"
for i, (k, v) in enumerate(hw_data):
    tbl.rows[i].cells[0].text = k
    tbl.rows[i].cells[1].text = v
    for c in tbl.rows[i].cells:
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c.paragraphs[0].runs[0].font.size = Pt(10.5)
        c.paragraphs[0].runs[0].font.name = "Calibri"
        set_cell_border(c)
    if i == 0:
        for c in tbl.rows[i].cells:
            set_cell_bg(c, "1A5376")
            tbl.rows[i].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            tbl.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
            tbl.rows[i].cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            tbl.rows[i].cells[1].paragraphs[0].runs[0].font.bold = True
    else:
        bg = "EAF4FB" if i % 2 == 0 else "FFFFFF"
        for c in tbl.rows[i].cells:
            set_cell_bg(c, bg)

doc.add_paragraph()

# ─── 4. IMPLEMENTATION ────────────────────────────────────────────────────────
heading(doc, "4. Implementation")

heading(doc, "4.1  CPU Implementation (Sequential 2D Array)", level=2)
body(doc, (
    "The CPU implementation stores the entire grid as a flat 1D array in row-major order (grid[y * width + x]). "
    "This is more cache-friendly than a true 2D pointer array. At each generation, a nested for-loop visits "
    "every cell, counts its 8 neighbors, and writes the result to a second (next) buffer. After all cells are "
    "processed, the two buffers are swapped."
))
bullet(doc, "Data type: uint8_t (1 byte per cell) to reduce memory bandwidth.")
bullet(doc, "Boundary: Toroidal wrap-around computed per-row to avoid repeated division.")
bullet(doc, "Timing: std::chrono::high_resolution_clock for microsecond accuracy.")

heading(doc, "4.2  CUDA GPU Implementation (Parallel)", level=2)
body(doc, (
    "The CUDA implementation maps the 2D grid directly onto a 2D layout of GPU threads. "
    "Each thread is responsible for exactly one cell. Two device buffers (d_grid and d_next_grid) "
    "are allocated in GPU VRAM. The grid stays on the GPU for all 100 iterations — no data is "
    "copied back to the CPU between steps. After 100 steps, the final result is copied back once."
))
bullet(doc, "Thread block size: 16 x 16 = 256 threads per block.")
bullet(doc, "Grid of blocks: ceil(N/16) x ceil(N/16) blocks to cover the full grid.")
bullet(doc, "Two CUDA kernels are implemented:")

p = doc.add_paragraph(style="List Bullet 2")
p.add_run("Kernel 1 — Global Memory:").bold = True
p.add_run(
    " Each thread reads its 8 neighbors directly from GPU VRAM (Global Memory). "
    "Threads are organized to access memory in a coalesced (sequential) pattern for maximum bandwidth."
)

p = doc.add_paragraph(style="List Bullet 2")
p.add_run("Kernel 2 — Shared Memory (Optimization):").bold = True
p.add_run(
    " Each block of 16x16 threads cooperatively loads an 18x18 tile (including 1-cell halo borders) "
    "into fast on-chip Shared Memory before computing. This reduces repeated reads from slow VRAM."
)

bullet(doc, "GPU timing: CUDA Events (cudaEventRecord) for precise kernel-only timing.")

# ─── 5. HIGH PERFORMANCE TECHNIQUES ──────────────────────────────────────────
heading(doc, "5. High-Performance Optimization Techniques")
body(doc, "The following techniques were applied to achieve maximum performance:")

p = doc.add_paragraph(style="List Number")
p.add_run("Ping-Pong VRAM Double Buffering: ").bold = True
p.add_run(
    "Both the current state (d_grid) and next state (d_next_grid) are kept in GPU memory for the entire "
    "100-step loop. Only pointer addresses are swapped between steps using std::swap(d_grid, d_next_grid). "
    "This eliminates PCIe bus transfer overhead between every generation."
)

p = doc.add_paragraph(style="List Number")
p.add_run("Coalesced Memory Access: ").bold = True
p.add_run(
    "The 2D block layout ensures that adjacent GPU threads (threadIdx.x = 0, 1, 2, ...) access "
    "adjacent memory addresses in the same row, allowing the GPU hardware to combine them into "
    "a single wide memory transaction."
)

p = doc.add_paragraph(style="List Number")
p.add_run("1-Byte Cell Representation (uint8_t): ").bold = True
p.add_run(
    "Each cell uses only 1 byte instead of 4 bytes (int). For a 4096x4096 grid, this saves "
    "48 MB of memory bandwidth per read pass, greatly improving GPU cache hit rates."
)

p = doc.add_paragraph(style="List Number")
p.add_run("Compiler Optimization (-O3): ").bold = True
p.add_run(
    "Both the CPU and CUDA code are compiled with maximum optimization, enabling auto-vectorization "
    "and loop unrolling on the CPU."
)

p = doc.add_paragraph(style="List Number")
p.add_run("Shared Memory Tiling: ").bold = True
p.add_run(
    "Each 16x16 thread block preloads its tile plus halo cells into __shared__ memory. "
    "Since shared memory is ~100x faster than global memory, this reduces latency for "
    "neighbor lookups."
)

doc.add_paragraph()

# ─── 6. BENCHMARK RESULTS ─────────────────────────────────────────────────────
heading(doc, "6. Benchmark Results (100 Iterations)")
body(doc, (
    "The following table shows the measured execution time in milliseconds for 100 generations "
    "on the CPU and two CUDA kernels, across six grid sizes."
))

# Table header + rows
headers = ["Grid Size", "Total Cells", "CPU Time (ms)", "CUDA Global (ms)", "CUDA Shared (ms)", "Speedup (Global)", "Speedup (Shared)"]
tbl2 = doc.add_table(rows=1 + len(rows), cols=len(headers))
tbl2.style = "Table Grid"

# Header row
for j, h in enumerate(headers):
    cell = tbl2.rows[0].cells[j]
    cell.text = h
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].runs[0]
    run.font.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = "Calibri"
    set_cell_bg(cell, "1F6E5E")
    set_cell_border(cell)

# Data rows
for i, r in enumerate(rows):
    row_cells = tbl2.rows[i + 1].cells
    values = [
        f"{r['size']}x{r['size']}",
        f"{r['cells']:,}",
        f"{r['cpu']:.2f}",
        f"{r['glob']:.2f}",
        f"{r['shrd']:.2f}",
        f"{r['sp_glob']:.2f}x",
        f"{r['sp_shrd']:.2f}x",
    ]
    bg = "F0FAF7" if i % 2 == 0 else "FFFFFF"
    for j, val in enumerate(values):
        cell = row_cells[j]
        cell.text = val
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].runs[0]
        run.font.size = Pt(9.5)
        run.font.name = "Calibri"
        set_cell_bg(cell, bg)
        set_cell_border(cell)

doc.add_paragraph()

# ─── 7. CHARTS ────────────────────────────────────────────────────────────────
heading(doc, "7. Performance Charts")

if os.path.exists("hpc_benchmark_summary.png"):
    body(doc, "Figure 1: Combined benchmark summary — Execution Time (left) and Speedup Factor (right).")
    doc.add_picture("hpc_benchmark_summary.png", width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

if os.path.exists("plot_execution_time.png"):
    body(doc, "Figure 2: Execution time comparison — CPU rises sharply while CUDA stays near the bottom.")
    doc.add_picture("plot_execution_time.png", width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

if os.path.exists("plot_speedup.png"):
    body(doc, "Figure 3: GPU Speedup multiplier. As grid size grows, the GPU advantage increases steadily.")
    doc.add_picture("plot_speedup.png", width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

# ─── 8. DISCUSSION ────────────────────────────────────────────────────────────
heading(doc, "8. Discussion & Analysis")

# Pull real numbers from last row for discussion
last = rows[-1]
first = rows[0]
body(doc, (
    f"For the smallest grid (128 x 128), the CPU completed 100 generations in {first['cpu']:.2f} ms "
    f"while CUDA (Global) took {first['glob']:.2f} ms, giving a speedup of {first['sp_glob']:.2f}x. "
    f"The speedup is modest here because the grid has only {first['cells']:,} cells, "
    "which is not enough to fully saturate the GPU's parallel cores."
))

body(doc, (
    f"For the largest grid tested (4096 x 4096), the CPU took {last['cpu']:.2f} ms "
    f"({last['cpu']/1000:.2f} seconds) while CUDA (Global) took only {last['glob']:.2f} ms. "
    f"This gives a speedup of {last['sp_glob']:.2f}x — meaning the GPU is {last['sp_glob']:.1f} "
    f"times faster than the CPU for {last['cells']:,} cells."
))

body(doc, (
    "The Shared Memory kernel is consistently slower than the Global Memory kernel in this case. "
    "This is because Conway's Game of Life is a simple, low-arithmetic-intensity problem: each cell "
    "only needs to add 8 neighbors. The overhead of loading halo tiles into shared memory adds latency "
    "that outweighs the benefit for this particular computation pattern. "
    "For heavier computations (e.g. image filtering or fluid simulation), shared memory would typically win."
))

body(doc, (
    "The GPU speedup grows steadily from 15x at 128x128 to 32x at 4096x4096. "
    "This scaling behavior confirms that GPU parallelism becomes more effective as the problem size grows."
))

# ─── 9. CONCLUSION ────────────────────────────────────────────────────────────
heading(doc, "9. Conclusion")
body(doc, (
    "We successfully implemented Conway's Game of Life using both a sequential CPU method (2D arrays) "
    "and a parallel GPU method (CUDA). The CUDA implementation using coalesced global memory access "
    "and ping-pong VRAM double buffering delivered the best performance, achieving up to "
    f"{last['sp_glob']:.1f}x speedup over the CPU at a grid size of 4096 x 4096."
))
body(doc, (
    "Key performance techniques used:"
))
bullet(doc, "Keeping data on GPU for all 100 iterations (no per-step CPU-GPU copying).")
bullet(doc, "Coalesced global memory access with 16x16 thread blocks.")
bullet(doc, "1-byte cell representation to reduce memory bandwidth by 4x.")
bullet(doc, "Maximum compiler optimization with -O3 flag.")

body(doc, (
    "The interactive GUI application allows real-time switching between CPU and CUDA modes, "
    "live grid drawing, preset patterns (Gosper Glider Gun), and on-demand 100-iteration benchmarking "
    "directly from the dashboard."
))

# ─── Save ─────────────────────────────────────────────────────────────────────
out = "Conway_Game_of_Life_HPC_Report.docx"
doc.save(out)
print(f"[OK] Report saved: {out}")
